"""
WirelessAgent System Prompt Structure Visualization
Creates an editable PowerPoint slide showing the prompt design structure
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    
    colors = {
        'title': RGBColor(0x1F, 0x49, 0x7D),
        'overview': RGBColor(0x2E, 0x75, 0xB6),
        'embb': RGBColor(0x00, 0xB0, 0x50),
        'urllc': RGBColor(0xFF, 0x00, 0x00),
        'mmtc': RGBColor(0xFF, 0xC0, 0x00),
        'core': RGBColor(0x70, 0x30, 0xA0),
        'flow': RGBColor(0x00, 0x70, 0xC0),
        'tech': RGBColor(0x83, 0x4C, 0x2B),
        'error': RGBColor(0xC0, 0x00, 0x00),
        'test': RGBColor(0x00, 0xB0, 0xF0),
    }
    
    title = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.3), Inches(0.2),
        Inches(12.7), Inches(0.6)
    )
    title.fill.solid()
    title.fill.fore_color.rgb = colors['title']
    title.line.fill.background()
    tf = title.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "WirelessAgent System Prompt Structure"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    
    overview = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(4.5), Inches(0.95),
        Inches(4.3), Inches(0.7)
    )
    overview.fill.solid()
    overview.fill.fore_color.rgb = colors['overview']
    overview.line.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    tf = overview.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Project Overview\n5G/6G Network Slicing Management"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    
    core_cap_title = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.3), Inches(1.8),
        Inches(2.2), Inches(0.4)
    )
    core_cap_title.fill.solid()
    core_cap_title.fill.fore_color.rgb = colors['core']
    core_cap_title.line.fill.background()
    tf = core_cap_title.text_frame
    p = tf.paragraphs[0]
    p.text = "Core Capabilities"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    
    slice_types_title = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.3), Inches(2.3),
        Inches(2.2), Inches(0.35)
    )
    slice_types_title.fill.solid()
    slice_types_title.fill.fore_color.rgb = RGBColor(0x50, 0x50, 0x50)
    slice_types_title.line.fill.background()
    tf = slice_types_title.text_frame
    p = tf.paragraphs[0]
    p.text = "Slice Types"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    
    embb_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.3), Inches(2.75),
        Inches(2.2), Inches(1.1)
    )
    embb_box.fill.solid()
    embb_box.fill.fore_color.rgb = colors['embb']
    embb_box.line.color.rgb = RGBColor(0x00, 0x80, 0x40)
    tf = embb_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "eMBB\nEnhanced Mobile Broadband"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    p = tf.add_paragraph()
    p.text = "90 MHz | 40-50ms\nVideo, VR/AR, Gaming"
    p.font.size = Pt(8)
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    
    urllc_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.3), Inches(3.95),
        Inches(2.2), Inches(1.1)
    )
    urllc_box.fill.solid()
    urllc_box.fill.fore_color.rgb = colors['urllc']
    urllc_box.line.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    tf = urllc_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "URLLC\nUltra-Reliable Low-Latency"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    p = tf.add_paragraph()
    p.text = "30 MHz | <5ms\nVehicle, Surgery, Control"
    p.font.size = Pt(8)
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    
    mmtc_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.3), Inches(5.15),
        Inches(2.2), Inches(1.1)
    )
    mmtc_box.fill.solid()
    mmtc_box.fill.fore_color.rgb = colors['mmtc']
    mmtc_box.line.color.rgb = RGBColor(0xD4, 0xA0, 0x00)
    tf = mmtc_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "mMTC\nMassive Machine-Type"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p.alignment = PP_ALIGN.CENTER
    p = tf.add_paragraph()
    p.text = "10 MHz | Variable\nIoT, Sensors, Meters"
    p.font.size = Pt(8)
    p.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p.alignment = PP_ALIGN.CENTER
    
    cqi_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(2.7), Inches(2.3),
        Inches(2.4), Inches(1.9)
    )
    cqi_box.fill.solid()
    cqi_box.fill.fore_color.rgb = RGBColor(0x00, 0x70, 0xC0)
    cqi_box.line.color.rgb = RGBColor(0x00, 0x50, 0x90)
    tf = cqi_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "CQI-Based Calculation"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    p = tf.add_paragraph()
    p.text = "CQI Range: 1-15"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    p = tf.add_paragraph()
    p.text = "Low (1-3): Reduced BW"
    p.font.size = Pt(8)
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.LEFT
    p = tf.add_paragraph()
    p.text = "Medium (4-8): Standard"
    p.font.size = Pt(8)
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.LEFT
    p = tf.add_paragraph()
    p.text = "Good (9-12): Higher"
    p.font.size = Pt(8)
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.LEFT
    p = tf.add_paragraph()
    p.text = "Excellent (13-15): Max"
    p.font.size = Pt(8)
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.LEFT
    
    formula_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(2.7), Inches(4.35),
        Inches(2.4), Inches(0.7)
    )
    formula_box.fill.solid()
    formula_box.fill.fore_color.rgb = RGBColor(0x40, 0x40, 0x40)
    formula_box.line.color.rgb = RGBColor(0x60, 0x60, 0x60)
    tf = formula_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Formula: BW = Max_BW × (CQI/15)"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x00, 0xFF, 0x00)
    p.alignment = PP_ALIGN.CENTER
    
    ray_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(2.7), Inches(5.2),
        Inches(2.4), Inches(1.0)
    )
    ray_box.fill.solid()
    ray_box.fill.fore_color.rgb = RGBColor(0x70, 0x30, 0xA0)
    ray_box.line.color.rgb = RGBColor(0x50, 0x20, 0x80)
    tf = ray_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Ray Tracing Simulation"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    p = tf.add_paragraph()
    p.text = "SNR, RX Power, LOS/NLOS"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    
    flow_title = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(5.3), Inches(1.8),
        Inches(7.7), Inches(0.4)
    )
    flow_title.fill.solid()
    flow_title.fill.fore_color.rgb = colors['flow']
    flow_title.line.fill.background()
    tf = flow_title.text_frame
    p = tf.paragraphs[0]
    p.text = "Data Flow Pipeline"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    
    flow_steps = [
        ("1. Load Ray\nTracing Data", (0x00, 0x70, 0xC0)),
        ("2. Generate\nUser Requests", (0x00, 0x90, 0xD0)),
        ("3. Intent\nClassification", (0x70, 0x30, 0xA0)),
        ("4. Slice Type\nDetermination", (0x00, 0xB0, 0x50)),
        ("5. Bandwidth\nCalculation", (0xFF, 0xC0, 0x00)),
        ("6. Resource\nTracking", (0xFF, 0x66, 0x00)),
        ("7. Accuracy\nValidation", (0x00, 0xB0, 0xF0)),
    ]
    
    start_x = 5.35
    box_width = 1.0
    box_height = 0.8
    gap = 0.08
    y_pos = 2.35
    
    for i, (text, color_tuple) in enumerate(flow_steps):
        x = start_x + i * (box_width + gap)
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x), Inches(y_pos),
            Inches(box_width), Inches(box_height)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(*color_tuple)
        box.line.color.rgb = RGBColor(
            max(0, color_tuple[0] - 40),
            max(0, color_tuple[1] - 40),
            max(0, color_tuple[2] - 40)
        )
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.alignment = PP_ALIGN.CENTER
        tf.paragraphs[0].space_before = Pt(4)
        
        if i < len(flow_steps) - 1:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Inches(x + box_width), Inches(y_pos + box_height/2 - 0.1),
                Inches(gap), Inches(0.2)
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = RGBColor(0x80, 0x80, 0x80)
            arrow.line.fill.background()
    
    tech_title = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(5.3), Inches(3.3),
        Inches(3.7), Inches(0.35)
    )
    tech_title.fill.solid()
    tech_title.fill.fore_color.rgb = colors['tech']
    tech_title.line.fill.background()
    tf = tech_title.text_frame
    p = tf.paragraphs[0]
    p.text = "Technical Specifications"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    
    tech_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(5.3), Inches(3.7),
        Inches(3.7), Inches(1.3)
    )
    tech_box.fill.solid()
    tech_box.fill.fore_color.rgb = RGBColor(0xF5, 0xE6, 0xD3)
    tech_box.line.color.rgb = colors['tech']
    tf = tech_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Frequency: 2.4 GHz"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p.alignment = PP_ALIGN.LEFT
    p = tf.add_paragraph()
    p.text = "Bandwidth: 20 MHz"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p.alignment = PP_ALIGN.LEFT
    p = tf.add_paragraph()
    p.text = "TX Power: 30 dBm"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p.alignment = PP_ALIGN.LEFT
    p = tf.add_paragraph()
    p.text = "RX Height: 1.5m"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p.alignment = PP_ALIGN.LEFT
    p = tf.add_paragraph()
    p.text = "Noise Figure: 8 dB"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p.alignment = PP_ALIGN.LEFT
    
    error_title = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(9.2), Inches(3.3),
        Inches(3.8), Inches(0.35)
    )
    error_title.fill.solid()
    error_title.fill.fore_color.rgb = colors['error']
    error_title.line.fill.background()
    tf = error_title.text_frame
    p = tf.paragraphs[0]
    p.text = "Error Handling"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    
    error_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(9.2), Inches(3.7),
        Inches(3.8), Inches(1.3)
    )
    error_box.fill.solid()
    error_box.fill.fore_color.rgb = RGBColor(0xFF, 0xE0, 0xE0)
    error_box.line.color.rgb = colors['error']
    tf = error_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "LLM API: Retry + Backoff"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p.alignment = PP_ALIGN.LEFT
    p = tf.add_paragraph()
    p.text = "RAG Failure: Fallback to LLM"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p.alignment = PP_ALIGN.LEFT
    p = tf.add_paragraph()
    p.text = "Invalid CQI: Clamp to [1,15]"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p.alignment = PP_ALIGN.LEFT
    p = tf.add_paragraph()
    p.text = "Missing Files: Clear Error"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p.alignment = PP_ALIGN.LEFT
    
    test_title = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(5.3), Inches(5.15),
        Inches(3.7), Inches(0.35)
    )
    test_title.fill.solid()
    test_title.fill.fore_color.rgb = colors['test']
    test_title.line.fill.background()
    tf = test_title.text_frame
    p = tf.paragraphs[0]
    p.text = "Testing Approach"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    
    test_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(5.3), Inches(5.55),
        Inches(3.7), Inches(1.0)
    )
    test_box.fill.solid()
    test_box.fill.fore_color.rgb = RGBColor(0xE0, 0xF5, 0xFF)
    test_box.line.color.rgb = colors['test']
    tf = test_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Intent Classification Accuracy"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p.alignment = PP_ALIGN.LEFT
    p = tf.add_paragraph()
    p.text = "CQI-Bandwidth Mapping Validation"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p.alignment = PP_ALIGN.LEFT
    p = tf.add_paragraph()
    p.text = "Token Usage Comparison"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p.alignment = PP_ALIGN.LEFT
    p = tf.add_paragraph()
    p.text = "Multi-Scenario Batch Testing"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p.alignment = PP_ALIGN.LEFT
    
    metrics_title = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(9.2), Inches(5.15),
        Inches(3.8), Inches(0.35)
    )
    metrics_title.fill.solid()
    metrics_title.fill.fore_color.rgb = RGBColor(0x00, 0x80, 0x80)
    metrics_title.line.fill.background()
    tf = metrics_title.text_frame
    p = tf.paragraphs[0]
    p.text = "Accuracy Metrics"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    
    metrics_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(9.2), Inches(5.55),
        Inches(3.8), Inches(1.0)
    )
    metrics_box.fill.solid()
    metrics_box.fill.fore_color.rgb = RGBColor(0xE0, 0xFF, 0xF0)
    metrics_box.line.color.rgb = RGBColor(0x00, 0x80, 0x80)
    tf = metrics_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Intent Understanding Accuracy"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p.alignment = PP_ALIGN.LEFT
    p = tf.add_paragraph()
    p.text = "Token Efficiency Tracking"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p.alignment = PP_ALIGN.LEFT
    p = tf.add_paragraph()
    p.text = "Resource Utilization Efficiency"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p.alignment = PP_ALIGN.LEFT
    
    keywords_title = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.3), Inches(6.4),
        Inches(12.7), Inches(0.35)
    )
    keywords_title.fill.solid()
    keywords_title.fill.fore_color.rgb = RGBColor(0x50, 0x50, 0x50)
    keywords_title.line.fill.background()
    tf = keywords_title.text_frame
    p = tf.paragraphs[0]
    p.text = "Intent Classification Keywords"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    
    kw_embb = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.3), Inches(6.85),
        Inches(4.2), Inches(0.5)
    )
    kw_embb.fill.solid()
    kw_embb.fill.fore_color.rgb = RGBColor(0xE0, 0xF5, 0xE0)
    kw_embb.line.color.rgb = colors['embb']
    tf = kw_embb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "eMBB: video, 4K, streaming, VR, AR, gaming, download, conference"
    p.font.size = Pt(8)
    p.font.color.rgb = RGBColor(0x00, 0x60, 0x30)
    p.alignment = PP_ALIGN.CENTER
    
    kw_urllc = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(4.6), Inches(6.85),
        Inches(4.2), Inches(0.5)
    )
    kw_urllc.fill.solid()
    kw_urllc.fill.fore_color.rgb = RGBColor(0xFF, 0xE0, 0xE0)
    kw_urllc.line.color.rgb = colors['urllc']
    tf = kw_urllc.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "URLLC: control, autonomous, remote, surgery, vehicle, real-time"
    p.font.size = Pt(8)
    p.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    p.alignment = PP_ALIGN.CENTER
    
    kw_mmtc = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(8.9), Inches(6.85),
        Inches(4.1), Inches(0.5)
    )
    kw_mmtc.fill.solid()
    kw_mmtc.fill.fore_color.rgb = RGBColor(0xFF, 0xF5, 0xE0)
    kw_mmtc.line.color.rgb = colors['mmtc']
    tf = kw_mmtc.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "mMTC: sensor, meter, IoT, monitor, smart, wearable, tracking"
    p.font.size = Pt(8)
    p.font.color.rgb = RGBColor(0x80, 0x60, 0x00)
    p.alignment = PP_ALIGN.CENTER
    
    return prs

if __name__ == "__main__":
    prs = create_presentation()
    output_path = "f:/code/wirelessagent/WirelessAgent_Prompt_Structure.pptx"
    prs.save(output_path)
    print(f"PowerPoint saved to: {output_path}")
