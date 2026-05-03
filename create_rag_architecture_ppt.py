from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

prs = Presentation()
prs.slide_width = Inches(16)
prs.slide_height = Inches(9)

slide_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(slide_layout)

def add_title(slide, text, left, top, width, height, font_size=28, bold=True, color=RGBColor(0x1a, 0x3c, 0x6e)):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xf0, 0xf4, 0xf8)
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = PP_ALIGN.CENTER
    tf.margin_left = Pt(5)
    tf.margin_right = Pt(5)
    tf.margin_top = Pt(3)
    tf.margin_bottom = Pt(3)
    return shape

def add_component_box(slide, text, left, top, width, height, fill_color, border_color, font_size=11, bold=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(4)
    tf.margin_right = Pt(4)
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = RGBColor(0x2d, 0x3e, 0x50)
    p.alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return shape

def add_data_box(slide, text, left, top, width, height, fill_color, border_color, font_size=9):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(3)
    tf.margin_right = Pt(3)
    tf.margin_top = Pt(1)
    tf.margin_bottom = Pt(1)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = RGBColor(0x3d, 0x4f, 0x5f)
    p.alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return shape

def add_arrow(slide, start_x, start_y, end_x, end_y, color=RGBColor(0x5a, 0x7a, 0x9a)):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        start_x, start_y, end_x, end_y
    )
    connector.line.color.rgb = color
    connector.line.width = Pt(2)
    return connector

title_shape = add_title(
    slide, "RAG历史记忆检索优化架构图",
    Inches(0.3), Inches(0.2), Inches(15.4), Inches(0.5),
    font_size=24, bold=True, color=RGBColor(0x1a, 0x3c, 0x6e)
)

input_box = add_component_box(
    slide, "用户请求\n(User Request)\n+ CQI + Location",
    Inches(0.4), Inches(1.2), Inches(1.8), Inches(0.9),
    RGBColor(0xe8, 0xf4, 0xfd), RGBColor(0x21, 0x96, 0xf3),
    font_size=11, bold=True
)

fusion_box = add_component_box(
    slide, "FusionRetriever\n融合检索器",
    Inches(2.8), Inches(1.35), Inches(2.2), Inches(0.8),
    RGBColor(0xe3, 0xf2, 0xfd), RGBColor(0x19, 0x76, 0xd2),
    font_size=12, bold=True
)

hybrid_box = add_component_box(
    slide, "HybridRetriever\n混合检索器",
    Inches(5.6), Inches(0.7), Inches(2.0), Inches(0.8),
    RGBColor(0xf3, 0xe5, 0xf5), RGBColor(0x7b, 0x1f, 0xa2),
    font_size=11, bold=True
)

historical_box = add_component_box(
    slide, "HistoricalRetriever\n历史记忆检索器",
    Inches(5.6), Inches(1.8), Inches(2.0), Inches(0.8),
    RGBColor(0xe8, 0xf5, 0xe9), RGBColor(0x38, 0x8e, 0x3c),
    font_size=11, bold=True
)

vector_store = add_component_box(
    slide, "向量存储\n(Vector Store)",
    Inches(8.2), Inches(0.4), Inches(1.5), Inches(0.6),
    RGBColor(0xfc, 0xe4, 0xec), RGBColor(0xc2, 0x18, 0x5b),
    font_size=9
)

bm25_store = add_component_box(
    slide, "BM25索引\n(BM25 Index)",
    Inches(8.2), Inches(1.1), Inches(1.5), Inches(0.6),
    RGBColor(0xfc, 0xe4, 0xec), RGBColor(0xc2, 0x18, 0x5b),
    font_size=9
)

kb_box = add_component_box(
    slide, "知识库\n(Knowledge Base)",
    Inches(10.0), Inches(0.4), Inches(1.6), Inches(1.3),
    RGBColor(0xfb, 0xe9, 0xe7), RGBColor(0xd3, 0x2f, 0x2f),
    font_size=10, bold=True
)

memory_store = add_component_box(
    slide, "HistoricalMemoryStore\n历史记忆存储",
    Inches(8.2), Inches(1.9), Inches(2.0), Inches(0.9),
    RGBColor(0xe0, 0xf2, 0xf1), RGBColor(0x00, 0x69, 0x6a),
    font_size=10, bold=True
)

cases_file = add_data_box(
    slide, "cases.jsonl\n案例文件",
    Inches(10.4), Inches(1.75), Inches(1.2), Inches(0.45),
    RGBColor(0xf5, 0xf5, 0xf5), RGBColor(0x75, 0x75, 0x75),
    font_size=8
)

faiss_index = add_data_box(
    slide, "cases.index\nFAISS索引",
    Inches(10.4), Inches(2.3), Inches(1.2), Inches(0.45),
    RGBColor(0xf5, 0xf5, 0xf5), RGBColor(0x75, 0x75, 0x75),
    font_size=8
)

result_box = add_component_box(
    slide, "融合检索结果\n{kb_context, historical_context,\nconfidence, sources}",
    Inches(2.8), Inches(2.5), Inches(2.2), Inches(0.8),
    RGBColor(0xff, 0xf8, 0xe1), RGBColor(0xff, 0xa0, 0x00),
    font_size=9
)

agent_box = add_component_box(
    slide, "Agent工作流\n(WA_DS_V3_FKB)",
    Inches(5.6), Inches(3.2), Inches(2.4), Inches(2.0),
    RGBColor(0xe0, 0xf7, 0xfa), RGBColor(0x00, 0x97, 0xa7),
    font_size=11, bold=True
)

intent_node = add_data_box(
    slide, "understand_intent\n意图理解节点\n(调用LLM)",
    Inches(5.8), Inches(3.4), Inches(2.0), Inches(0.65),
    RGBColor(0xb2, 0xeb, 0xf2), RGBColor(0x00, 0x97, 0xa7),
    font_size=8
)

allocate_node = add_data_box(
    slide, "allocate_slice_type\n切片分配节点",
    Inches(5.8), Inches(4.15), Inches(2.0), Inches(0.45),
    RGBColor(0xb2, 0xeb, 0xf2), RGBColor(0x00, 0x97, 0xa7),
    font_size=8
)

record_node = add_data_box(
    slide, "record_allocation_case\n案例记录节点",
    Inches(5.8), Inches(4.7), Inches(2.0), Inches(0.45),
    RGBColor(0xc8, 0xe6, 0xc9), RGBColor(0x38, 0x8e, 0x3c),
    font_size=8
)

llm_box = add_component_box(
    slide, "LLM\n大语言模型",
    Inches(8.5), Inches(3.5), Inches(1.6), Inches(0.7),
    RGBColor(0xe1, 0xbe, 0xe7), RGBColor(0x8e, 0x24, 0xaa),
    font_size=10, bold=True
)

output_box = add_component_box(
    slide, "分配结果\n(Slice Allocation)\n切片类型 + 带宽",
    Inches(0.4), Inches(3.5), Inches(1.8), Inches(0.9),
    RGBColor(0xd4, 0xed, 0xda), RGBColor(0x28, 0xa7, 0x45),
    font_size=10, bold=True
)

add_arrow(slide, Inches(2.2), Inches(1.65), Inches(2.8), Inches(1.7))
add_arrow(slide, Inches(5.0), Inches(1.55), Inches(5.6), Inches(1.1))
add_arrow(slide, Inches(5.0), Inches(1.9), Inches(5.6), Inches(2.2))
add_arrow(slide, Inches(7.6), Inches(1.0), Inches(8.2), Inches(0.7))
add_arrow(slide, Inches(7.6), Inches(1.3), Inches(8.2), Inches(1.4))
add_arrow(slide, Inches(7.6), Inches(2.2), Inches(8.2), Inches(2.35))
add_arrow(slide, Inches(9.7), Inches(0.7), Inches(10.0), Inches(0.9))
add_arrow(slide, Inches(9.7), Inches(1.4), Inches(10.0), Inches(1.2))
add_arrow(slide, Inches(10.2), Inches(2.0), Inches(10.4), Inches(2.0))
add_arrow(slide, Inches(10.2), Inches(2.35), Inches(10.4), Inches(2.55))
add_arrow(slide, Inches(3.9), Inches(2.15), Inches(3.9), Inches(2.5))
add_arrow(slide, Inches(5.0), Inches(2.9), Inches(5.6), Inches(3.6))
add_arrow(slide, Inches(7.8), Inches(3.72), Inches(8.5), Inches(3.85))
add_arrow(slide, Inches(6.8), Inches(4.05), Inches(6.8), Inches(4.15))
add_arrow(slide, Inches(6.8), Inches(4.6), Inches(6.8), Inches(4.7))
add_arrow(slide, Inches(9.2), Inches(2.8), Inches(9.2), Inches(3.5))
add_arrow(slide, Inches(5.6), Inches(4.9), Inches(5.0), Inches(4.9))
add_arrow(slide, Inches(4.2), Inches(4.9), Inches(4.2), Inches(4.4))
add_arrow(slide, Inches(2.8), Inches(3.9), Inches(2.2), Inches(3.95))

legend_y = Inches(5.6)
legend_items = [
    ("输入/输出", RGBColor(0xe8, 0xf4, 0xfd), RGBColor(0x21, 0x96, 0xf3)),
    ("检索器", RGBColor(0xe3, 0xf2, 0xfd), RGBColor(0x19, 0x76, 0xd2)),
    ("存储/索引", RGBColor(0xfc, 0xe4, 0xec), RGBColor(0xc2, 0x18, 0x5b)),
    ("历史记忆", RGBColor(0xe0, 0xf2, 0xf1), RGBColor(0x00, 0x69, 0x6a)),
    ("LLM/Agent", RGBColor(0xe1, 0xbe, 0xe7), RGBColor(0x8e, 0x24, 0xaa)),
]

for i, (label, fill, border) in enumerate(legend_items):
    x = Inches(0.4 + i * 3.1)
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, legend_y, Inches(0.4), Inches(0.25))
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = border
    box.line.width = Pt(1)
    
    text_box = slide.shapes.add_textbox(x + Inches(0.5), legend_y, Inches(2.4), Inches(0.25))
    tf = text_box.text_frame
    p = tf.paragraphs[0]
    p.text = label
    p.font.size = Pt(10)
    p.font.color.rgb = RGBColor(0x3d, 0x4f, 0x5f)

flow_title = slide.shapes.add_textbox(Inches(0.4), Inches(6.0), Inches(15.2), Inches(0.3))
tf = flow_title.text_frame
p = tf.paragraphs[0]
p.text = "数据流向: 用户请求 → FusionRetriever → [HybridRetriever → 知识库] + [HistoricalRetriever → 历史存储] → 融合结果 → understand_intent(调用LLM) → allocate_slice_type → 分配结果"
p.font.size = Pt(9)
p.font.color.rgb = RGBColor(0x5a, 0x6a, 0x7a)
p.alignment = PP_ALIGN.LEFT

feedback_title = slide.shapes.add_textbox(Inches(0.4), Inches(6.35), Inches(15.2), Inches(0.25))
tf = feedback_title.text_frame
p = tf.paragraphs[0]
p.text = "自学习闭环: record_allocation_case → HistoricalMemoryStore (持久化案例，丰富历史库)"
p.font.size = Pt(9)
p.font.color.rgb = RGBColor(0x00, 0x69, 0x6a)
p.font.bold = True
p.alignment = PP_ALIGN.LEFT

benefits_title = slide.shapes.add_textbox(Inches(0.4), Inches(6.65), Inches(15.2), Inches(0.25))
tf = benefits_title.text_frame
p = tf.paragraphs[0]
p.text = "预期收益: Token消耗 ~15000→~8000/请求 | 检索准确率: 关键词匹配→语义+历史案例 | 解决冷启动问题"
p.font.size = Pt(9)
p.font.color.rgb = RGBColor(0x2e, 0x7d, 0x32)
p.font.bold = True
p.alignment = PP_ALIGN.LEFT

files_title = slide.shapes.add_textbox(Inches(0.4), Inches(6.95), Inches(15.2), Inches(0.25))
tf = files_title.text_frame
p = tf.paragraphs[0]
p.text = "关键文件: historical_memory_store.py (新增) | rag_system.py (修改: HistoricalRetriever, FusionRetriever) | WA_DS_V3_FKB.py (修改: 集成融合检索)"
p.font.size = Pt(9)
p.font.color.rgb = RGBColor(0x15, 0x65, 0xc0)
p.alignment = PP_ALIGN.LEFT

prs.save('F:/code/wirelessagent/docs/RAG历史记忆检索优化架构图.pptx')
print("PPT架构图已更新: F:/code/wirelessagent/docs/RAG历史记忆检索优化架构图.pptx")
